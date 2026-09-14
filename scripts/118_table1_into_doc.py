"""Rebuild Table 1 IN PLACE inside the hand-edited Results docx (surgical: only the table is
replaced; all other text is preserved). Reads /tmp/table1.csv (script 112). Full row set organised
in blocks: burden · clinical profile · mode of detection · treatment outcome · place characteristics.
"""
import pandas as pd
from docx import Document
from docx.shared import Pt

REP=("/Users/evelynlepkadelima/Library/CloudStorage/GoogleDrive-evelynlepka@gmail.com/"
     "My Drive/WHO modelling Project/SP-TB-spatial-analyses/Reports/SP_TB_Manuscript_Results.docx")
t=pd.read_csv("/tmp/table1.csv").set_index("group")
GROUPS=["Metropolitan — hotspot","Metropolitan — non-hotspot","Interior — hotspot",
        "Interior — non-hotspot","Favela regions (subset)","State total"]
HEAD=["","Metropolitan\nhotspot","Metropolitan\nnon-hotspot","Interior\nhotspot",
      "Interior\nnon-hotspot","Favela\nregions*","State\ntotal"]

def i(v): return f"{int(round(v)):,}"
def f1(v): return f"{v:.1f}"
def sgn(v): return f"{v:+.2f}"
def val(k,fmt): return lambda g: fmt(t.loc[g,k])
def npct(nk,pk): return lambda g: f"{int(round(t.loc[g,nk])):,} ({t.loc[g,pk]:.1f})"

ROWS=[
 ("row","No. of regions",                    val("n_reg",i)),
 ("row","Adult population, n (%)",            npct("pop","pop_pct")),
 ("row","Incident TB cases, n (%)",           npct("cases","cases_pct")),
 ("row","TB incidence, /100 000/year",        val("inc",f1)),
 ("row","TB mortality, /100 000/year",        val("mort",f1)),
 ("sub","Clinical profile",None),
 ("row","Pulmonary TB, %",                    val("pulm",f1)),
 ("row","HIV-positive, % of those tested",    val("hiv_pos",f1)),
 ("sub","Mode of detection, %",None),
 ("row","Outpatient (patient demand)",        val("outpt",f1)),
 ("row","Emergency / urgency presentation",   val("er",f1)),
 ("row","Diagnosed during hospitalisation",   val("hospdx",f1)),
 ("row","Active case finding",                val("acf",f1)),
 ("row","Contact investigation",              val("contact",f1)),
 ("sub","Treatment outcome, %",None),
 ("row","Treatment success (cure)",           val("cure",f1)),
 ("row","Loss to follow-up",                  val("aband_pct",f1)),
 ("row","Death from TB",                       val("tbdeath",f1)),
 ("sub","Place characteristics",None),
 ("row","Household income, R$ (mean)",        val("income",i)),
 ("row","Adult illiteracy, %",                val("illit",f1)),
 ("row","Residents per household",            val("residents",f1)),
 ("row","Favela population, %",               val("favela_pct",f1)),
 ("row","Vulnerability index (z-score)",      val("vuln",sgn)),
]

doc=Document(REP)
assert len(doc.tables)==1, f"expected 1 table, found {len(doc.tables)}"
tbl=doc.tables[0]
for row in list(tbl.rows): tbl._tbl.remove(row._tr)   # clear, keep style + column grid

hdr=tbl.add_row().cells
for j,h in enumerate(HEAD):
    hdr[j].text=h
    for p in hdr[j].paragraphs:
        for r in p.runs: r.font.bold=True; r.font.size=Pt(8.5)
for kind,lab,fn in ROWS:
    cells=tbl.add_row().cells
    if kind=="sub":
        a=cells[0]
        for c in cells[1:]: a=a.merge(c)
        a.text=lab
        for p in a.paragraphs:
            for r in p.runs: r.font.bold=True; r.font.italic=True; r.font.size=Pt(8.5)
    else:
        cells[0].text=lab
        for r in cells[0].paragraphs[0].runs: r.font.size=Pt(9)
        for j,g in enumerate(GROUPS):
            cells[j+1].text=fn(g)
            for p in cells[j+1].paragraphs:
                for r in p.runs: r.font.size=Pt(9)
doc.save(REP)
print("Table 1 rebuilt:",len(tbl.rows),"rows x",len(tbl.columns),"cols")
